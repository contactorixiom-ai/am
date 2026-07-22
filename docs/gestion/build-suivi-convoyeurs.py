# -*- coding: utf-8 -*-
"""Classeur de gestion des convoyages Axis Import — France / Europe (V1)."""
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

NAVY = "0B2545"
GOLD = "C9A55C"
GOLD_SOFT = "F3ECD9"
INK = "1A2230"
MUTED = "6B7280"
LINE = "D8DEE7"
INPUT_FILL = "FFF7E0"   # jaune clair : cellules à remplir
CALC_FILL = "EFF3F8"    # gris-bleu : cellules calculées
WHITE = "FFFFFF"

F = "Arial"
thin = Side(style="thin", color=LINE)
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(sz=10, bold=False, color=INK, italic=False):
    return Font(name=F, size=sz, bold=bold, color=color, italic=italic)

def fill(hexc):
    return PatternFill("solid", fgColor=hexc)

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left(wrap=False):
    return Alignment(horizontal="left", vertical="center", wrap_text=wrap)

wb = openpyxl.Workbook()

# ════════════════════════════════════════════════════════════════════
# Données sources
# ════════════════════════════════════════════════════════════════════
CONVOYEURS = [
    ("Mamadou Bah", "+33 6 12 34 56 78", "mamadou.bah@axis-import.fr", "Actif"),
    ("Karim Diallo", "+33 6 23 45 67 89", "karim.diallo@axis-import.fr", "Actif"),
    ("Awa Ndiaye", "+33 6 34 56 78 90", "awa.ndiaye@axis-import.fr", "Actif"),
    ("Julien Moreau", "+33 6 45 67 89 01", "julien.moreau@axis-import.fr", "Actif"),
]
CLIENTS = [
    ("Transports Diallo SARL", "Professionnel", "+33 1 40 00 00 00"),
    ("M. Amadou Sow", "Particulier", "+33 6 00 00 00 01"),
    ("Auto Prestige Lyon", "Professionnel", "+33 4 78 00 00 00"),
]
VEHICULES = [
    ("AB-123-CD", "Peugeot 508", "Berline"),
    ("EF-456-GH", "Renault Master", "Utilitaire"),
    ("GT-820-RK", "Renault T480", "Poids lourd"),
    ("IJ-789-KL", "BMW Série 3", "Berline"),
]
# Villes France + Europe (convoyage) : (ville, pays)
VILLES = [
    # France
    ("Paris","France"),("Marseille","France"),("Lyon","France"),("Toulouse","France"),
    ("Nice","France"),("Nantes","France"),("Montpellier","France"),("Strasbourg","France"),
    ("Bordeaux","France"),("Lille","France"),("Rennes","France"),("Reims","France"),
    ("Le Havre","France"),("Saint-Étienne","France"),("Toulon","France"),("Grenoble","France"),
    ("Dijon","France"),("Angers","France"),("Nîmes","France"),("Clermont-Ferrand","France"),
    ("Le Mans","France"),("Aix-en-Provence","France"),("Brest","France"),("Tours","France"),
    ("Amiens","France"),("Limoges","France"),("Annecy","France"),("Perpignan","France"),
    ("Besançon","France"),("Metz","France"),("Orléans","France"),("Rouen","France"),
    ("Mulhouse","France"),("Caen","France"),("Nancy","France"),("Avignon","France"),
    ("Poitiers","France"),("Dunkerque","France"),("Pau","France"),("Bayonne","France"),
    ("La Rochelle","France"),("Cannes","France"),("Calais","France"),("Colmar","France"),
    ("Chambéry","France"),("Valence","France"),("Troyes","France"),("Lorient","France"),
    # Belgique / Pays-Bas / Luxembourg
    ("Bruxelles","Belgique"),("Anvers","Belgique"),("Liège","Belgique"),("Gand","Belgique"),
    ("Charleroi","Belgique"),("Amsterdam","Pays-Bas"),("Rotterdam","Pays-Bas"),
    ("La Haye","Pays-Bas"),("Eindhoven","Pays-Bas"),("Luxembourg","Luxembourg"),
    # Allemagne
    ("Berlin","Allemagne"),("Munich","Allemagne"),("Francfort","Allemagne"),("Cologne","Allemagne"),
    ("Hambourg","Allemagne"),("Stuttgart","Allemagne"),("Düsseldorf","Allemagne"),
    ("Sarrebruck","Allemagne"),("Fribourg","Allemagne"),
    # Suisse / Autriche
    ("Genève","Suisse"),("Zurich","Suisse"),("Bâle","Suisse"),("Lausanne","Suisse"),
    ("Berne","Suisse"),("Vienne","Autriche"),("Salzbourg","Autriche"),
    # Italie
    ("Milan","Italie"),("Turin","Italie"),("Rome","Italie"),("Gênes","Italie"),
    ("Bologne","Italie"),("Florence","Italie"),("Vintimille","Italie"),
    # Espagne / Portugal
    ("Madrid","Espagne"),("Barcelone","Espagne"),("Bilbao","Espagne"),("Valence","Espagne"),
    ("Saragosse","Espagne"),("Séville","Espagne"),("Lisbonne","Portugal"),("Porto","Portugal"),
    # Royaume-Uni / Irlande
    ("Londres","Royaume-Uni"),("Manchester","Royaume-Uni"),("Birmingham","Royaume-Uni"),
    ("Dublin","Irlande"),
    # Europe du Nord / Est
    ("Copenhague","Danemark"),("Stockholm","Suède"),("Oslo","Norvège"),
    ("Prague","Tchéquie"),("Varsovie","Pologne"),("Cracovie","Pologne"),("Budapest","Hongrie"),
]

MAXROWS = 150  # lignes de saisie dans Convoyages

# ════════════════════════════════════════════════════════════════════
# Feuille ACCUEIL
# ════════════════════════════════════════════════════════════════════
acc = wb.active
acc.title = "Accueil"
acc.sheet_view.showGridLines = False
for c in "ABCDEFGH":
    acc.column_dimensions[c].width = 15
acc.column_dimensions["A"].width = 3

acc.merge_cells("B2:H3")
acc["B2"] = "AXIS IMPORT — GESTION DES CONVOYAGES"
acc["B2"].font = font(20, True, NAVY)
acc["B2"].alignment = left()
acc.merge_cells("B4:H4")
acc["B4"] = "Suivi d'activité des convoyeurs · France & Europe · Version 1 (base de données)"
acc["B4"].font = font(11, False, MUTED)

def section(cell, text):
    acc[cell] = text
    acc[cell].font = font(12, True, WHITE)
    acc[cell].fill = fill(NAVY)
    acc[cell].alignment = left()

acc.merge_cells("B6:H6"); section("B6", "  Contenu du classeur")
sheets_info = [
    ("Convoyages", "Saisie de chaque trajet : date, convoyeur, lieux, distance, montant, charges, TVA, bénéfice."),
    ("Synthèse", "Totaux automatiques par mois, par année et par convoyeur (nb, distance, CA, bénéfice)."),
    ("Convoyeurs", "Liste des convoyeurs — alimente le menu déroulant. Ajoute une ligne pour un nouveau convoyeur."),
    ("Clients", "Liste des clients (menu déroulant)."),
    ("Véhicules", "Liste des véhicules convoyés (menu déroulant)."),
    ("Lieux", "Villes de France & d'Europe — alimente les menus Départ / Arrivée."),
    ("Paramètres", "Taux de TVA et grille des options (modifiable au même endroit)."),
]
r = 7
for name, desc in sheets_info:
    acc[f"B{r}"] = name
    acc[f"B{r}"].font = font(10.5, True, NAVY)
    acc.merge_cells(f"C{r}:H{r}")
    acc[f"C{r}"] = desc
    acc[f"C{r}"].font = font(10, False, INK)
    acc[f"C{r}"].alignment = left(wrap=True)
    r += 1

r += 1
acc.merge_cells(f"B{r}:H{r}"); section(f"B{r}", "  Mode d'emploi")
r += 1
legend = [
    ("Cellules JAUNES", "à remplir par toi (date, convoyeur, lieux, distance, montant, carburant, péages, nourriture).", INPUT_FILL),
    ("Cellules GRISES", "calculées automatiquement (TVA collectée, TTC, TVA déductible, bénéfice/perte). Ne pas y écrire.", CALC_FILL),
    ("Menus déroulants", "convoyeur, client, véhicule, lieux : clique sur la cellule puis sur la flèche.", WHITE),
    ("Ajouter un convoyeur", "va dans l'onglet Convoyeurs et écris son nom sur une nouvelle ligne : il apparaît dans le menu.", WHITE),
    ("Date", "saisie au format JJ/MM/AAAA. (Le calendrier cliquable arrive en V3 avec les macros.)", WHITE),
]
for label, desc, fillc in legend:
    acc[f"B{r}"] = label
    acc[f"B{r}"].font = font(10, True, NAVY)
    acc[f"B{r}"].fill = fill(fillc)
    acc[f"B{r}"].border = border
    acc.merge_cells(f"C{r}:H{r}")
    acc[f"C{r}"] = desc
    acc[f"C{r}"].font = font(10, False, INK)
    acc[f"C{r}"].alignment = left(wrap=True)
    r += 1

r += 1
acc.merge_cells(f"B{r}:H{r}"); section(f"B{r}", "  Feuille de route")
r += 1
roadmap = [
    ("V1 — Base de données", "Sheets, listes déroulantes, calculs TVA/TTC/bénéfice, mise en forme.  (cette version)"),
    ("V2 — Tableau de bord", "Statistiques jour/mois/année, graphiques, classement des convoyeurs."),
    ("V3 — Formulaires VBA", "Boutons Ajouter/Modifier/Supprimer, recherche instantanée, calendrier intégré."),
    ("V4 — Facturation", "Facturation automatique, export PDF, numérotation, gestion TVA."),
]
for label, desc in roadmap:
    acc[f"B{r}"] = label
    acc[f"B{r}"].font = font(10, True, GOLD.replace("C9A55C","8A6D2B"))
    acc.merge_cells(f"C{r}:H{r}")
    acc[f"C{r}"] = desc
    acc[f"C{r}"].font = font(10, False, INK)
    acc[f"C{r}"].alignment = left(wrap=True)
    r += 1

# ════════════════════════════════════════════════════════════════════
# Feuilles sources (Convoyeurs, Clients, Véhicules, Lieux)
# ════════════════════════════════════════════════════════════════════
def build_list_sheet(title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1, title).font = font(14, True, NAVY)
    for j, h in enumerate(headers, 1):
        c = ws.cell(2, j, h)
        c.font = font(10, True, WHITE); c.fill = fill(NAVY)
        c.alignment = center(); c.border = border
        ws.column_dimensions[get_column_letter(j)].width = widths[j-1]
    for i, row in enumerate(rows, 3):
        for j, val in enumerate(row, 1):
            c = ws.cell(i, j, val)
            c.font = font(10, False, INK); c.border = border
            c.fill = fill(INPUT_FILL if j == 1 else WHITE)
            c.alignment = left()
    # lignes vides jusqu'à 200 pour permettre les ajouts (menus déroulants)
    for i in range(3 + len(rows), 201):
        c = ws.cell(i, 1); c.fill = fill(INPUT_FILL); c.border = border
    ws.freeze_panes = "A3"
    return ws

build_list_sheet("Convoyeurs", ["Nom du convoyeur", "Téléphone", "Email", "Statut"],
                 CONVOYEURS, [26, 18, 30, 12])
build_list_sheet("Clients", ["Nom du client", "Type", "Contact"],
                 CLIENTS, [30, 16, 20])
build_list_sheet("Véhicules", ["Immatriculation", "Marque / Modèle", "Catégorie"],
                 VEHICULES, [18, 24, 16])

# Lieux : col A ville, B pays, C libellé combiné (source des menus Départ/Arrivée)
lieux = wb.create_sheet("Lieux")
lieux.sheet_view.showGridLines = False
lieux.merge_cells("A1:C1")
lieux["A1"] = "Lieux — France & Europe"; lieux["A1"].font = font(14, True, NAVY)
for j, h in enumerate(["Ville", "Pays", "Libellé (menu Départ / Arrivée)"], 1):
    c = lieux.cell(2, j, h); c.font = font(10, True, WHITE); c.fill = fill(NAVY)
    c.alignment = center(); c.border = border
lieux.column_dimensions["A"].width = 22
lieux.column_dimensions["B"].width = 16
lieux.column_dimensions["C"].width = 30
villes_sorted = sorted(VILLES, key=lambda x: (x[1], x[0]))
for i, (ville, pays) in enumerate(villes_sorted, 3):
    lieux.cell(i, 1, ville).font = font(10, False, INK)
    lieux.cell(i, 2, pays).font = font(10, False, INK)
    lc = lieux.cell(i, 3, f"={get_column_letter(1)}{i}&\" · \"&{get_column_letter(2)}{i}")
    lc.font = font(10, False, INK)
    for j in range(1, 4):
        lieux.cell(i, j).border = border
        lieux.cell(i, j).alignment = left()
lieux.freeze_panes = "A3"
LIEUX_LAST = 2 + len(villes_sorted)

# ════════════════════════════════════════════════════════════════════
# Feuille PARAMÈTRES
# ════════════════════════════════════════════════════════════════════
par = wb.create_sheet("Paramètres")
par.sheet_view.showGridLines = False
par.column_dimensions["A"].width = 34
par.column_dimensions["B"].width = 18
par.column_dimensions["C"].width = 34
par.merge_cells("A1:C1")
par["A1"] = "Paramètres"; par["A1"].font = font(14, True, NAVY)
par["A3"] = "Taux de TVA (France)"; par["A3"].font = font(10, True, INK)
par["B3"] = 0.20; par["B3"].number_format = "0%"
par["B3"].fill = fill(INPUT_FILL); par["B3"].border = border; par["B3"].font = font(10, True, "0000FF")
par["C3"] = "Modifiable ici — utilisé par tout le classeur."; par["C3"].font = font(9, False, MUTED)
par["A4"] = "Devise"; par["A4"].font = font(10, True, INK)
par["B4"] = "EUR (€)"; par["B4"].fill = fill(INPUT_FILL); par["B4"].border = border; par["B4"].font = font(10, False, "0000FF")

par.merge_cells("A6:C6")
par["A6"] = "  Grille des options (rappel — aligné sur l'application)"
par["A6"].font = font(11, True, WHITE); par["A6"].fill = fill(NAVY)
opts = [
    ("Express 24h", "+15 %", "Supplément sur le sous-total"),
    ("Assurance Premium", "+29,61 €", "Plafond 350 000 €"),
    ("Porte-à-porte", "Gratuit", "Inclus"),
    ("Enlèvement weekend", "+60 €", "Forfait"),
    ("Récupération colis (domicile)", "25 €", "Forfait"),
    ("Enlèvement colis (point relais)", "5 €", "Forfait"),
]
par.cell(7,1,"Option").font = font(10, True, WHITE); par.cell(7,1).fill = fill(GOLD)
par.cell(7,2,"Tarif").font = font(10, True, WHITE); par.cell(7,2).fill = fill(GOLD)
par.cell(7,3,"Détail").font = font(10, True, WHITE); par.cell(7,3).fill = fill(GOLD)
for j in range(1,4):
    par.cell(7,j).alignment=center(); par.cell(7,j).border=border
for i,(a,b,c) in enumerate(opts, 8):
    par.cell(i,1,a).font=font(10,False,INK)
    par.cell(i,2,b).font=font(10,True,NAVY); par.cell(i,2).alignment=center()
    par.cell(i,3,c).font=font(10,False,MUTED)
    for j in range(1,4): par.cell(i,j).border=border

# ════════════════════════════════════════════════════════════════════
# Feuille CONVOYAGES (saisie principale)
# ════════════════════════════════════════════════════════════════════
cv = wb.create_sheet("Convoyages")
cv.sheet_view.showGridLines = False
HEAD_ROW = 3
DATA_START = 4
DATA_END = DATA_START + MAXROWS - 1

cv.merge_cells("A1:O1")
cv["A1"] = "CONVOYAGES — France / Europe"
cv["A1"].font = font(16, True, NAVY)
cv.merge_cells("A2:O2")
cv["A2"] = "Une ligne = un trajet.  Cellules jaunes = à remplir · cellules grises = calculées automatiquement."
cv["A2"].font = font(9.5, False, MUTED)

cols = [
    ("Date", 12, "input"),
    ("Convoyeur", 20, "input"),
    ("Client", 22, "input"),
    ("Véhicule", 16, "input"),
    ("Lieu Départ", 22, "input"),
    ("Lieu Arrivée", 22, "input"),
    ("Distance (km)", 12, "input"),
    ("Montant HT (€)", 14, "input"),
    ("TVA collectée (€)", 14, "calc"),
    ("Montant TTC (€)", 14, "calc"),
    ("Carburant (€)", 12, "input"),
    ("Péages (€)", 11, "input"),
    ("Nourriture (€)", 12, "input"),
    ("TVA déductible (€)", 15, "calc"),
    ("Bénéfice / Perte (€)", 16, "calc"),
]
for j,(h,w,kind) in enumerate(cols, 1):
    c = cv.cell(HEAD_ROW, j, h)
    c.font = font(9.5, True, WHITE)
    c.fill = fill(NAVY if kind=="input" else "23405f")
    c.alignment = center(wrap=True); c.border = border
    cv.column_dimensions[get_column_letter(j)].width = w
# colonnes helper Année / Mois (masquées)
cv.cell(HEAD_ROW, 16, "Année").font = font(9, True, MUTED)
cv.cell(HEAD_ROW, 17, "Mois").font = font(9, True, MUTED)
cv.column_dimensions["P"].hidden = True
cv.column_dimensions["Q"].hidden = True

euro = '#,##0.00 €'
eur_neg = '#,##0.00 €;[Red]-#,##0.00 €'
TVA = "Paramètres!$B$3"

for row in range(DATA_START, DATA_END + 1):
    # Colonnes de saisie (jaune)
    for j in (1,2,3,4,5,6,7,8,11,12,13):
        c = cv.cell(row, j); c.fill = fill(INPUT_FILL); c.border = border
        c.font = font(10, False, INK)
    cv.cell(row,1).number_format = "DD/MM/YYYY"
    for j in (7,8,11,12,13):
        cv.cell(row, j).number_format = euro if j != 7 else '#,##0 "km"'
    # Formules (gris)
    I = f"=IF($H{row}=\"\",\"\",$H{row}*{TVA})"
    J = f"=IF($H{row}=\"\",\"\",$H{row}+$I{row})"
    N = f"=IF(SUM($K{row}:$M{row})=0,\"\",SUM($K{row}:$M{row})*({TVA}/(1+{TVA})))"
    O = f"=IF($H{row}=\"\",\"\",$H{row}-SUM($K{row}:$M{row})+IF($N{row}=\"\",0,$N{row}))"
    P = f"=IF($A{row}=\"\",\"\",YEAR($A{row}))"
    Q = f"=IF($A{row}=\"\",\"\",MONTH($A{row}))"
    for col, form, numfmt in [(9,I,euro),(10,J,euro),(14,N,euro),(15,O,eur_neg)]:
        c = cv.cell(row, col, form)
        c.fill = fill(CALC_FILL); c.border = border
        c.font = font(10, False, INK); c.number_format = numfmt
    cv.cell(row,16,P).font = font(9, False, MUTED)
    cv.cell(row,17,Q).font = font(9, False, MUTED)

cv.freeze_panes = "A4"

# Ligne d'exemple (au-dessus, non dans la plage de saisie) — retirée ensuite ?
# On met un exemple sur la 1re ligne de données pour montrer le format.
example = [datetime.date(2026, 1, 15), "Mamadou Bah", "Transports Diallo SARL", "Peugeot 508",
           "Paris · France", "Lyon · France", 465, 520, None, None, 78, 42, 18]
for j, val in enumerate(example, 1):
    if val is not None:
        cv.cell(DATA_START, j, val)

# ── Menus déroulants (data validation) ───────────────────────────────
def add_dv(ws, formula, col_letter, r1, r2, allow_blank=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank, showDropDown=False)
    dv.error = "Choisis une valeur dans la liste (ou ajoute-la dans l'onglet source)."
    dv.errorTitle = "Valeur hors liste"
    dv.prompt = "Sélectionne dans le menu déroulant."
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{r1}:{col_letter}{r2}")

add_dv(cv, "=Convoyeurs!$A$3:$A$200", "B", DATA_START, DATA_END)
add_dv(cv, "=Clients!$A$3:$A$200", "C", DATA_START, DATA_END)
add_dv(cv, "=Véhicules!$A$3:$A$200", "D", DATA_START, DATA_END)
add_dv(cv, f"=Lieux!$C$3:$C${LIEUX_LAST}", "E", DATA_START, DATA_END)
add_dv(cv, f"=Lieux!$C$3:$C${LIEUX_LAST}", "F", DATA_START, DATA_END)

# Validation de date (calendrier visuel = V3/VBA)
dvd = DataValidation(type="date", operator="between",
                     formula1="DATE(2020,1,1)", formula2="DATE(2035,12,31)",
                     allow_blank=True)
dvd.error = "Saisis une date valide (JJ/MM/AAAA)."
dvd.errorTitle = "Date invalide"
dvd.prompt = "Saisis la date du convoyage (JJ/MM/AAAA)."
dvd.promptTitle = "Date"
cv.add_data_validation(dvd)
dvd.add(f"A{DATA_START}:A{DATA_END}")

# ════════════════════════════════════════════════════════════════════
# Feuille SYNTHÈSE (par mois / par année / par convoyeur)
# ════════════════════════════════════════════════════════════════════
sy = wb.create_sheet("Synthèse")
sy.sheet_view.showGridLines = False
for c,w in zip("ABCDEFGHIJ",[16,14,12,14,13,12,13,15,15,16]):
    sy.column_dimensions[c].width = w

sy.merge_cells("A1:J1")
sy["A1"] = "SYNTHÈSE D'ACTIVITÉ"; sy["A1"].font = font(16, True, NAVY)
sy["A3"] = "Année analysée :"; sy["A3"].font = font(11, True, INK); sy["A3"].alignment = Alignment(horizontal="right", vertical="center")
sy["B3"] = 2026; sy["B3"].font = font(12, True, "0000FF"); sy["B3"].fill = fill(INPUT_FILL)
sy["B3"].border = border; sy["B3"].alignment = center(); sy["B3"].number_format = "0"
sy["C3"] = "← change l'année ici"; sy["C3"].font = font(9, False, MUTED)

P = f"Convoyages!$P$4:$P${DATA_END}"   # Année
Q = f"Convoyages!$Q$4:$Q${DATA_END}"   # Mois
def RG(col):  # plage colonne données Convoyages
    return f"Convoyages!${col}$4:${col}${DATA_END}"

# ── Tableau par mois ────────────────────────────────────────────────
sy.merge_cells("A5:J5"); sy["A5"] = "  Par mois (année sélectionnée)"
sy["A5"].font = font(12, True, WHITE); sy["A5"].fill = fill(NAVY)
head = ["Mois","Nb convoyages","Distance (km)","CA HT (€)","Carburant (€)",
        "Péages (€)","Nourriture (€)","TVA collectée (€)","TVA déductible (€)","Bénéfice (€)"]
for j,h in enumerate(head,1):
    c = sy.cell(6,j,h); c.font=font(9.5,True,WHITE); c.fill=fill("23405f")
    c.alignment=center(wrap=True); c.border=border
mois_noms = ["Janvier","Février","Mars","Avril","Mai","Juin","Juillet","Août",
             "Septembre","Octobre","Novembre","Décembre"]
euro = '#,##0.00 €'; eur_neg='#,##0.00 €;[Red]-#,##0.00 €'
first_month_row = 7
for m in range(1,13):
    r = first_month_row + m - 1
    sy.cell(r,1,mois_noms[m-1]).font=font(10,False,INK)
    crit = f"{P},$B$3,{Q},{m}"
    sy.cell(r,2, f"=COUNTIFS({crit})").number_format="0"
    sy.cell(r,3, f"=SUMIFS({RG('G')},{crit})").number_format='#,##0'
    sy.cell(r,4, f"=SUMIFS({RG('H')},{crit})").number_format=euro
    sy.cell(r,5, f"=SUMIFS({RG('K')},{crit})").number_format=euro
    sy.cell(r,6, f"=SUMIFS({RG('L')},{crit})").number_format=euro
    sy.cell(r,7, f"=SUMIFS({RG('M')},{crit})").number_format=euro
    sy.cell(r,8, f"=SUMIFS({RG('I')},{crit})").number_format=euro
    sy.cell(r,9, f"=SUMIFS({RG('N')},{crit})").number_format=euro
    sy.cell(r,10,f"=SUMIFS({RG('O')},{crit})").number_format=eur_neg
    for j in range(1,11):
        sy.cell(r,j).border=border
        if j>1: sy.cell(r,j).font=font(10,False,INK)
# Total année
tr = first_month_row + 12
sy.cell(tr,1,"TOTAL " ).font=font(10,True,WHITE)
sy.cell(tr,1).value = "TOTAL ANNÉE"
sy.cell(tr,1).font=font(10,True,WHITE); sy.cell(tr,1).fill=fill(GOLD)
for j in range(2,11):
    col = get_column_letter(j)
    c = sy.cell(tr,j, f"=SUM({col}{first_month_row}:{col}{tr-1})")
    c.font=font(10,True,NAVY); c.fill=fill(GOLD_SOFT); c.border=border
    c.number_format = "0" if j==2 else ('#,##0' if j==3 else (eur_neg if j==10 else euro))
sy.cell(tr,1).border=border

# ── Classement par convoyeur (toutes années) ────────────────────────
base = tr + 2
sy.merge_cells(f"A{base}:J{base}")
sy.cell(base,1,"  Par convoyeur (toutes périodes)")
sy.cell(base,1).font=font(12,True,WHITE); sy.cell(base,1).fill=fill(NAVY)
ch = ["Convoyeur","Nb convoyages","Distance (km)","CA HT (€)","Carburant (€)",
      "Péages (€)","Nourriture (€)","TVA collectée (€)","TVA déductible (€)","Bénéfice (€)"]
hr = base+1
for j,h in enumerate(ch,1):
    c=sy.cell(hr,j,h); c.font=font(9.5,True,WHITE); c.fill=fill("23405f")
    c.alignment=center(wrap=True); c.border=border
for i,(nom,_,_,_) in enumerate(CONVOYEURS):
    r = hr+1+i
    ref = f"=Convoyeurs!$A${3+i}"
    sy.cell(r,1, ref).font=font(10,False,INK)
    crit = f"{RG('B')},$A{r}"
    sy.cell(r,2, f"=COUNTIFS({crit})").number_format="0"
    sy.cell(r,3, f"=SUMIFS({RG('G')},{crit})").number_format='#,##0'
    sy.cell(r,4, f"=SUMIFS({RG('H')},{crit})").number_format=euro
    sy.cell(r,5, f"=SUMIFS({RG('K')},{crit})").number_format=euro
    sy.cell(r,6, f"=SUMIFS({RG('L')},{crit})").number_format=euro
    sy.cell(r,7, f"=SUMIFS({RG('M')},{crit})").number_format=euro
    sy.cell(r,8, f"=SUMIFS({RG('I')},{crit})").number_format=euro
    sy.cell(r,9, f"=SUMIFS({RG('N')},{crit})").number_format=euro
    sy.cell(r,10,f"=SUMIFS({RG('O')},{crit})").number_format=eur_neg
    for j in range(1,11):
        sy.cell(r,j).border=border
        if j>1: sy.cell(r,j).font=font(10,False,INK)
sy.cell(hr+1+len(CONVOYEURS)+1,1,
        "Astuce : le détail par jour se lit directement dans l'onglet Convoyages (triable par Date).").font=font(9,False,MUTED)

sy.freeze_panes = "A7"

# ── ordre des onglets ────────────────────────────────────────────────
order = ["Accueil","Convoyages","Synthèse","Convoyeurs","Clients","Véhicules","Lieux","Paramètres"]
wb._sheets.sort(key=lambda s: order.index(s.title))

# Force Excel/LibreOffice à tout recalculer à l'ouverture (aucune valeur mise en
# cache par openpyxl — le tableur calcule dès l'ouverture).
wb.calculation.fullCalcOnLoad = True

out = "/tmp/claude-0/-home-user-am/600e012e-ba72-5d4f-b8a2-e9bfbc0d47af/scratchpad/Suivi-Convoyeurs-Axis-V1.xlsx"
wb.save(out)
print("saved", out)
